import json
import tempfile
import unittest
from pathlib import Path

import openpyxl
import weekly_report as report


class ReportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def frame(self, rows):
        p = self.root / 'raw.json'
        p.write_text(json.dumps([{'results': rows}]))
        return report.load_rows(str(p))

    def test_site_isolation_and_no_old_domain_alias(self):
        rows = [dict(site=s, visitor_id='same', path='/products', day='2026-10-30', duration=8)
                for s in ['https://WWW.VCARPETS.COM/', 'vishomecarpet.com', 'inflatcustom.com', '']]
        frame = self.frame(rows)
        self.assertEqual(set(frame.site), {'vcarpets.com', 'vishomecarpet.com', 'inflatcustom.com', 'unknown'})
        self.assertEqual(len(report.build_visitor_frame(frame)), 4)
        self.assertEqual(len(report.build_page_frame(frame)), 4)

    def test_bad_durations_and_site(self):
        frame = self.frame([dict(duration=x, site='https://user:pass@example.com') for x in ['NaN', -8, 'inf', 10**20, 8]])
        self.assertEqual(frame.duration.tolist(), [0, 0, 0, 86400, 8])
        self.assertEqual(set(frame.site), {'unknown'})

    def test_aggregate_excludes_identity_and_suppresses_small_cells(self):
        rows = [dict(site='vcarpets.com', visitor_id=str(i), path='/products/name@example.com',
                     day='2026-10-30', traffic_channel='direct', duration=8,
                     org='PRIVATE_ORG', domain='PRIVATE_DOMAIN', ip_hash='PRIVATE_HASH',
                     session_id='PRIVATE_SESSION', query='PRIVATE_QUERY', gclid='PRIVATE_GCLID') for i in range(5)]
        rows.append(dict(site='inflatcustom.com', visitor_id='one', path='/products', day='2026-10-30'))
        frame = self.frame(rows)
        aggregate = report.build_aggregate_frame(frame)
        self.assertEqual(len(aggregate), 1)
        self.assertEqual(aggregate.iloc[0]['Page'], '/products')
        self.assertEqual(aggregate.iloc[0]['Observed Sampling Seconds'], 40)
        self.assertNotIn('PRIVATE', aggregate.to_json())
        output = self.root / 'aggregate.xlsx'
        report.main(str(self.root/'raw.json'), str(output), True)
        workbook = openpyxl.load_workbook(output)
        self.assertEqual(workbook.sheetnames, ['Read Me', 'Site Aggregates'])
        self.assertNotIn('PRIVATE', str([[c.value for row in sheet for c in row] for sheet in workbook]))

    def test_excel_formula_injection_is_text(self):
        self.frame([dict(site='vcarpets.com', visitor_id='id', org='=HYPERLINK("x")', path='/contact', country=' @SUM(1)', duration=8)])
        output = self.root / 'detail.xlsx'
        report.main(str(self.root/'raw.json'), str(output))
        workbook = openpyxl.load_workbook(output)
        self.assertFalse(any(c.data_type == 'f' for sheet in workbook for row in sheet for c in row))

    def test_site_filter_and_empty_export(self):
        self.frame([dict(site='vishomecarpet.com', visitor_id='id'), dict(site='vcarpets.com', visitor_id='id')])
        output = self.root/'filter.xlsx'
        report.main(str(self.root/'raw.json'), str(output), site='vcarpets.com')
        workbook = openpyxl.load_workbook(output)
        self.assertEqual(workbook['Site Summary']['A2'].value, 'vcarpets.com')
        self.assertEqual(workbook['Site Summary'].max_row, 2)
        self.frame([])
        report.main(str(self.root/'raw.json'), str(self.root/'empty.xlsx'), True)
        with self.assertRaises(ValueError):
            report.main(str(self.root/'raw.json'), str(self.root/'bad.xlsx'), True, 1)


if __name__ == '__main__':
    unittest.main()
